"""Generate native-anchor fixtures and measure LibreOffice PDF bookmark export.

Needs python-docx, openpyxl, python-pptx, pypdf and pdfplumber. Generated Office
and PDF files belong in an ignored reports/local directory, not the fixture tree.
"""

from __future__ import annotations

import argparse
import json
import re
import subprocess
from pathlib import Path


def generate(cases_path: Path, output: Path) -> dict:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt
    from openpyxl import Workbook
    from openpyxl.styles import Alignment
    from openpyxl.worksheet.pagebreak import Break
    from pptx import Presentation
    from pptx.util import Inches as SlideInches
    from pptx.util import Pt as SlidePt

    cases = json.loads(cases_path.read_text())
    output.mkdir(parents=True, exist_ok=True)
    documents = []
    doc = Document()
    doc.sections[0].page_width = Inches(6)
    doc.sections[0].page_height = Inches(8)
    doc.sections[0].left_margin = doc.sections[0].right_margin = Inches(0.7)
    doc.styles["Normal"].font.name = "Liberation Sans"
    doc.styles["Normal"].font.size = Pt(12)
    objects = []
    for index, case in enumerate(cases["docx"]):
        if "second_page" in case["cases"]:
            doc.add_page_break()
        if "table" in case["cases"]:
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Measurement"
            table.cell(0, 1).text = "Procedure"
            table.cell(1, 0).text = "Sediment"
            paragraph = table.cell(1, 1).paragraphs[0]
            anchor = {"kind": "table_cell", "table_index": 0, "row": 1, "column": 1}
        else:
            paragraph = doc.add_paragraph()
            anchor = {"kind": "paragraph", "paragraph_index": len(doc.paragraphs) - 1}
        bookmark = "capy" + case["id"].replace("_", "")
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(index))
        start.set(qn("w:name"), bookmark)
        paragraph._p.append(start)
        paragraph.add_run(case["text"])
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(index))
        paragraph._p.append(end)
        objects.append({**case, "anchor": anchor, "bookmark": bookmark})
    path = output / "wetland.docx"
    doc.save(path)
    documents.append(
        {
            "id": "wetland_docx",
            "format": "docx",
            "path": str(path.resolve()),
            "objects": objects,
        }
    )

    wb = Workbook()
    sheet = wb.active
    sheet.title = "Observations"
    sheet.append([cases["xlsx"]["header"], "Depth (cm)", "Action"])
    for row in range(2, 35):
        sheet.append(
            [
                f"Morning survey {row - 1}",
                row * 3,
                cases["xlsx"]["duplicate"]
                if row in (4, 24)
                else f"Inspect station {row - 1}",
            ]
        )
    sheet["A7"] = cases["xlsx"]["unicode"]
    sheet["C9"] = cases["xlsx"]["detail"]
    sheet["C9"].alignment = Alignment(wrap_text=True)
    sheet.row_dimensions[9].height = 50
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 14
    sheet.column_dimensions["C"].width = 34
    sheet.print_title_rows = "1:1"
    sheet.print_area = "A1:C34"
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.row_breaks.append(Break(id=18))
    secondary = wb.create_sheet("Control")
    secondary["A1"] = "Control site observations"
    secondary["B3"] = cases["xlsx"]["duplicate"]
    objects = []
    for sheet_name, cell, tags in [
        ("Observations", "A1", ["repeated_print_title"]),
        ("Observations", "A7", ["unicode"]),
        ("Observations", "C9", ["wrapped"]),
        ("Observations", "C4", ["duplicate"]),
        ("Observations", "C24", ["duplicate", "second_page"]),
        ("Control", "B3", ["duplicate", "other_sheet"]),
    ]:
        objects.append(
            {
                "id": sheet_name + "_" + cell,
                "text": wb[sheet_name][cell].value,
                "anchor": {"kind": "cell", "sheet": sheet_name, "cell": cell},
                "cases": tags,
            }
        )
    path = output / "wetland.xlsx"
    wb.save(path)
    documents.append(
        {
            "id": "wetland_xlsx",
            "format": "xlsx",
            "path": str(path.resolve()),
            "objects": objects,
        }
    )

    deck = Presentation()
    deck.slide_width, deck.slide_height = SlideInches(10), SlideInches(5.625)
    objects = []
    for slide_index in range(2):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        entries = [
            ("duplicate", 0.8, 0.7, 8, 0.8),
            ("wrapped", 0.8, 1.8, 4.3, 1.5),
            ("unicode", 0.8, 4, 8, 0.7),
        ]
        for label, x, y, width, height in entries:
            shape = slide.shapes.add_textbox(
                SlideInches(x), SlideInches(y), SlideInches(width), SlideInches(height)
            )
            shape.text_frame.word_wrap = True
            shape.text = cases["pptx"][label]
            for paragraph in shape.text_frame.paragraphs:
                paragraph.font.name = "Liberation Sans"
                paragraph.font.size = SlidePt(20)
            objects.append(
                {
                    "id": f"slide{slide_index}_{label}",
                    "text": shape.text,
                    "anchor": {
                        "kind": "shape",
                        "slide_index": slide_index,
                        "shape_id": shape.shape_id,
                    },
                    "cases": [label],
                    "bounds_emu": [shape.left, shape.top, shape.width, shape.height],
                }
            )
        if slide_index == 1:
            shape = slide.shapes.add_textbox(
                SlideInches(5.4), SlideInches(2.4), SlideInches(3.4), SlideInches(0.6)
            )
            shape.text = cases["pptx"]["rotated"]
            shape.rotation = 18
            objects.append(
                {
                    "id": "rotated",
                    "text": shape.text,
                    "anchor": {
                        "kind": "shape",
                        "slide_index": 1,
                        "shape_id": shape.shape_id,
                    },
                    "cases": ["rotated"],
                    "bounds_emu": [shape.left, shape.top, shape.width, shape.height],
                }
            )
            group = slide.shapes.add_group_shape()
            shape = group.shapes.add_textbox(
                SlideInches(5.2), SlideInches(3.4), SlideInches(4), SlideInches(0.6)
            )
            shape.text = cases["pptx"]["grouped"]
            objects.append(
                {
                    "id": "grouped",
                    "text": shape.text,
                    "anchor": {
                        "kind": "shape",
                        "slide_index": 1,
                        "shape_id": shape.shape_id,
                        "group_id": group.shape_id,
                    },
                    "cases": ["grouped"],
                    "bounds_emu": [shape.left, shape.top, shape.width, shape.height],
                }
            )
    path = output / "wetland.pptx"
    deck.save(path)
    documents.append(
        {
            "id": "wetland_pptx",
            "format": "pptx",
            "path": str(path.resolve()),
            "objects": objects,
        }
    )
    manifest = {
        "schema": 1,
        "provenance": "Synthetic author-created benchmark cases; indices are zero-based except spreadsheet addresses.",
        "documents": documents,
    }
    (output / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n"
    )
    return manifest


def probe(manifest: dict, output: Path, soffice: str) -> None:
    import pdfplumber
    from pypdf import PdfReader

    source = Path(manifest["documents"][0]["path"])
    results = []
    for variant, options in [
        ("default", {}),
        (
            "named",
            {"ExportBookmarksToPDFDestination": {"type": "boolean", "value": "true"}},
        ),
        (
            "named_tagged",
            {
                "ExportBookmarksToPDFDestination": {"type": "boolean", "value": "true"},
                "UseTaggedPDF": {"type": "boolean", "value": "true"},
            },
        ),
    ]:
        target = output / variant
        target.mkdir(parents=True, exist_ok=True)
        profile = (target / "profile").resolve().as_uri()
        filter_name = "pdf:writer_pdf_Export" + (
            ":" + json.dumps(options) if options else ""
        )
        command = [
            soffice,
            "-env:UserInstallation=" + profile,
            "--headless",
            "--convert-to",
            filter_name,
            "--outdir",
            str(target.resolve()),
            str(source),
        ]
        converted = subprocess.run(
            command, capture_output=True, text=True, timeout=90, check=False
        )
        pdf = target / "wetland.pdf"
        if converted.returncode or not pdf.exists():
            raise RuntimeError(
                f"{variant} conversion failed: {converted.stdout} {converted.stderr}"
            )
        reader = PdfReader(pdf)
        anchors = []
        with pdfplumber.open(pdf) as parsed:
            page_text = [page.extract_text() for page in parsed.pages]
            for obj in manifest["documents"][0]["objects"]:
                destination = reader.named_destinations.get(obj["bookmark"])
                if destination is None:
                    destination = reader.named_destinations.get("/" + obj["bookmark"])
                if destination is None:
                    anchors.append({"id": obj["id"], "found": False})
                    continue
                page_index = reader.get_destination_page_number(destination)
                page = parsed.pages[page_index]
                phrase = " ".join(obj["text"].split()[:5])
                hits = page.search(
                    r"\s+".join(re.escape(word) for word in phrase.split()), regex=True
                )
                x, top = float(destination.left), page.height - float(destination.top)
                distance = min(
                    (abs(x - hit["x0"]) + abs(top - hit["top"]) for hit in hits),
                    default=None,
                )
                full_text_found = " ".join(obj["text"].split()) in " ".join(
                    (page.extract_text() or "").split()
                )
                anchors.append(
                    {
                        "id": obj["id"],
                        "found": True,
                        "page": page_index + 1,
                        "destination": {
                            "x": x,
                            "top": top,
                            "type": str(destination.typ),
                        },
                        "prefix": phrase,
                        "prefix_hits_on_page": [
                            {key: hit[key] for key in ("x0", "top", "x1", "bottom")}
                            for hit in hits
                        ],
                        "nearest_prefix_start_manhattan_pt": distance,
                        "full_text_found_on_page": full_text_found,
                    }
                )
        result = {
            "variant": variant,
            "command": command,
            "stdout": converted.stdout,
            "page_count": len(reader.pages),
            "tagged": "/StructTreeRoot" in reader.trailer["/Root"],
            "named_destination_count": len(reader.named_destinations),
            "anchors": anchors,
            "page_text": page_text,
        }
        (target / "result.json").write_text(
            json.dumps(result, ensure_ascii=False, indent=2) + "\n"
        )
        results.append(result)
    (output / "bookmark-results.json").write_text(
        json.dumps(results, ensure_ascii=False, indent=2) + "\n"
    )
    assert not any(anchor["found"] for anchor in results[0]["anchors"])
    for result in results[1:]:
        assert all(anchor["found"] for anchor in result["anchors"])
        assert all(
            anchor["nearest_prefix_start_manhattan_pt"] is not None
            and anchor["nearest_prefix_start_manhattan_pt"] < 20
            for anchor in result["anchors"]
        )
    assert results[2]["tagged"]
    print(
        json.dumps(
            [
                {
                    key: result[key]
                    for key in (
                        "variant",
                        "page_count",
                        "tagged",
                        "named_destination_count",
                    )
                }
                for result in results
            ]
        )
    )


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument(
        "--soffice", help="Explicit LibreOffice executable; omit to generate only."
    )
    parser.add_argument(
        "--cases",
        type=Path,
        default=Path(__file__).resolve().parents[1]
        / "fixtures/native-citations/cases.json",
    )
    args = parser.parse_args()
    manifest = generate(args.cases, args.output / "corpus")
    print(args.output / "corpus/manifest.json", flush=True)
    if args.soffice:
        probe(manifest, args.output, args.soffice)


if __name__ == "__main__":
    main()
