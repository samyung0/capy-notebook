"""Regenerate the small synthetic UAT inputs; never contact a provider."""

from datetime import datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches as SlideInches
from reportlab.pdfgen.canvas import Canvas

ROOT = Path(__file__).parent / "basic"
MARKER = "UAT_RUN_MARKER"
FACT = "Wetland plants absorb carbon and protect the shoreline."


def stable_zip(path):
    with ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, data in sorted(entries.items()):
            info = ZipInfo(name, (2026, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            archive.writestr(info, data)


def main():
    ROOT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.core_properties.created = doc.core_properties.modified = datetime(2026, 1, 1)  # noqa: DTZ001 - fixed OOXML metadata uses a naive timestamp
    doc.add_paragraph("Owner sentence: The launch code is LARCH-17.")
    doc.add_paragraph("Collaborator sentence: The survey starts in June.")
    doc.add_paragraph(FACT)
    doc.add_paragraph(MARKER)
    table = doc.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, (("Habitat", "Count"), ("Wetland", "42"))):
        for cell, value in zip(row.cells, values):
            cell.text = value
    doc.sections[0].header.paragraphs[0].text = "Capy synthetic field notes"
    doc.save(ROOT / "lesson.docx")

    workbook = Workbook()
    workbook.properties.created = workbook.properties.modified = datetime(2026, 1, 1)  # noqa: DTZ001 - fixed OOXML metadata uses a naive timestamp
    sheet = workbook.active
    sheet.title = "Grades"
    sheet.append(["Student", "Score", "Double score"])
    sheet.append(["Owner", 42, "=B2*2"])
    sheet.append(["Collaborator", 17])
    sheet.append([FACT])
    sheet.append([MARKER])
    sheet.column_dimensions["A"].width = 68
    sheet.column_dimensions["B"].width = 14
    sheet.column_dimensions["C"].width = 18
    sheet.print_options.horizontalCentered = True
    workbook.create_sheet("Notes")["A1"] = "Preserve this independent worksheet."
    workbook.save(ROOT / "grades.xlsx")

    deck = Presentation()
    deck.core_properties.created = deck.core_properties.modified = datetime(2026, 1, 1)  # noqa: DTZ001 - fixed OOXML metadata uses a naive timestamp
    for text in (
        "Owner sentence: The launch code is LARCH-17.",
        "Collaborator sentence: The survey starts in June.",
    ):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(
            SlideInches(1), SlideInches(1), SlideInches(8), SlideInches(1)
        ).text = text
        slide.shapes.add_textbox(
            SlideInches(1), SlideInches(3), SlideInches(8), SlideInches(1)
        ).text = FACT
    deck.slides[0].shapes.add_textbox(
        SlideInches(1), SlideInches(5), SlideInches(8), SlideInches(0.5)
    ).text = MARKER
    deck.save(ROOT / "lesson.pptx")

    buffer = BytesIO()
    pdf = Canvas(buffer, pagesize=(612, 792), invariant=1)
    pdf.setTitle("Capy synthetic field notes")
    for y, text in ((720, FACT), (695, "The launch code is LARCH-17."), (670, MARKER)):
        pdf.drawString(54, y, text)
    pdf.save()
    (ROOT / "digital.pdf").write_bytes(buffer.getvalue())
    (ROOT / "notes.txt").write_text(f"{FACT}\nThe launch code is LARCH-17.\n{MARKER}\n")
    (ROOT / "grades.csv").write_text(
        f'name,score,note\nOwner,42,"{FACT}"\nCollaborator,17,{MARKER}\n'
    )
    (ROOT / "delimiter-limit.csv").write_text(MARKER + "," * 100_000 + "\n")
    for name in ("lesson.docx", "grades.xlsx", "lesson.pptx"):
        stable_zip(ROOT / name)


if __name__ == "__main__":
    main()
