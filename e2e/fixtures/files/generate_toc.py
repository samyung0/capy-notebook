"""Regenerate toc/report.docx: a Word-style table of contents; never contacts a provider."""

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from generate_basic import FACT, MARKER, stable_zip

ROOT = Path(__file__).parent / "toc"
# (level, heading, bookmark): Word's TOC \o "1-3" \h \z \u over these headings.
HEADINGS = (
    (1, "Field survey", "_Toc100000001"),
    (2, "Sampling sites", "_Toc100000002"),
    (1, "Findings", "_Toc100000003"),
)


def run(xml):
    return f"<w:r>{xml}</w:r>"


def field(instruction, result):
    """A complex field as Word writes it: begin, instruction, separate, result, end."""
    return (
        run('<w:fldChar w:fldCharType="begin"/>')
        + run(f'<w:instrText xml:space="preserve"> {instruction} </w:instrText>')
        + run('<w:fldChar w:fldCharType="separate"/>')
        + result
        + run('<w:fldChar w:fldCharType="end"/>')
    )


def entry(level, text, bookmark, first):
    # The TOC field opens in the first entry's paragraph; each entry is a
    # hyperlink to its heading's bookmark holding a PAGEREF field.
    start = (
        run('<w:fldChar w:fldCharType="begin"/>')
        + run(
            '<w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText>'
        )
        + run('<w:fldChar w:fldCharType="separate"/>')
        if first
        else ""
    )
    link = (
        f'<w:hyperlink w:anchor="{bookmark}" w:history="1">'
        + run(f"<w:t>{text}</w:t>")
        + run("<w:tab/>")
        + field(f"PAGEREF {bookmark} \\h", run("<w:t>1</w:t>"))
        + "</w:hyperlink>"
    )
    return parse_xml(
        f'<w:p {nsdecls("w")}><w:pPr><w:pStyle w:val="TOC{level}"/>'
        '<w:tabs><w:tab w:val="right" w:leader="dot" w:pos="8630"/></w:tabs>'
        f"</w:pPr>{start}{link}</w:p>"
    )


def main():
    ROOT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.core_properties.created = doc.core_properties.modified = datetime(2026, 1, 1)  # noqa: DTZ001 - fixed OOXML metadata uses a naive timestamp
    for level in (1, 2):
        style = doc.styles.add_style(f"toc {level}", WD_STYLE_TYPE.PARAGRAPH)
        style.element.styleId = f"TOC{level}"
    doc.add_paragraph("Wetland field report", style="Title")
    doc.add_paragraph("Contents", style="TOC Heading")
    body = doc.element.body
    sect = body[-1]
    for index, (level, text, bookmark) in enumerate(HEADINGS):
        sect.addprevious(entry(level, text, bookmark, index == 0))
    end = run('<w:fldChar w:fldCharType="end"/>')
    sect.addprevious(parse_xml(f"<w:p {nsdecls('w')}>{end}</w:p>"))
    paragraphs = {
        "Field survey": FACT,
        "Sampling sites": "Three sampling sites line the eastern shore.",
        "Findings": MARKER,
    }
    for index, (level, text, bookmark) in enumerate(HEADINGS):
        heading = doc.add_heading(level=level)
        heading._p.append(
            parse_xml(
                f'<w:bookmarkStart {nsdecls("w")} w:id="{index}" w:name="{bookmark}"/>'
            )
        )
        heading.add_run(text)
        heading._p.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="{index}"/>'))
        doc.add_paragraph(paragraphs[text])
    doc.save(ROOT / "report.docx")
    stable_zip(ROOT / "report.docx")


if __name__ == "__main__":
    main()
